"""Przelicza kalkulator w LibreOffice i porównuje z niezależnym obliczeniem w Pythonie.

    python3 tests/test_kalkulator.py
"""
import math
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "products"))
from build_kalkulator import DEFAULT_INPUTS, MIN_JOB, RATES, STORY_MULT, build  # noqa: E402

SCENARIOS = {
    "pokazowy (dom + podjazd + rynny, 2 piętra)": {},
    "mały podjazd, poniżej minimum": dict(driveway=300, house=0, gutters=0, tier="LOW",
                                          miles=20, chemicals=15, hours=2.5),
    "dach soft wash, HIGH, 1 piętro": dict(driveway=0, house=0, roof=2000, gutters=0,
                                           stories=1, tier="HIGH", hours=10, chemicals=120),
    "same rynny, tanio i daleko": dict(driveway=0, house=0, gutters=180, stories=1, tier="LOW",
                                       miles=60, hours=4, chemicals=0),
}


def expected(v):
    col = ["LOW", "MID", "HIGH"].index(v["tier"]) + 3
    rate = {r[0]: r[col] for r in RATES}
    mult = dict(STORY_MULT)[v["stories"]]
    lines = [round(rate["driveway"] * v["driveway"], 2),
             round(rate["house"] * v["house"] * mult, 2),
             round(rate["roof"] * v["roof"], 2),
             round(rate["gutters"] * v["gutters"] * mult, 2)]
    sub = sum(lines)
    price = 0 if sub == 0 else round(max(sub, MIN_JOB))
    cost = v["hours"] * v["labor"] + v["miles"] * v["per_mile"] + v["chemicals"]
    margin = 0 if price == 0 else (price - cost) / price
    need = math.ceil(cost / (1 - v["target"]) / 10) * 10
    if price == 0:
        verdict = "Enter"
    elif margin >= v["target"]:
        verdict = "OK"
    elif margin >= v["target"] - 0.1:
        verdict = "LOW"
    else:
        verdict = "REJECT"
    return price, cost, margin, need, verdict


def recalc(src, outdir):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir",
                    str(outdir), str(src)], check=True, capture_output=True, timeout=180)
    return load_workbook(Path(outdir) / Path(src).name, data_only=True)


def main():
    ok = True
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for i, (name, over) in enumerate(SCENARIOS.items()):
            v = {**DEFAULT_INPUTS, **over}
            src = build(tmp / f"s{i}.xlsx", v)
            wb = recalc(src, tmp / "out")
            s = wb["START"]
            got = (s["B21"].value, s["B22"].value, s["B24"].value, s["B25"].value, s["B27"].value)
            exp = expected(v)
            errors = [c.coordinate for ws in wb for row in ws.iter_rows() for c in row
                      if isinstance(c.value, str) and c.value.startswith("#")]
            match = (abs(got[0] - exp[0]) < 0.01 and abs(got[1] - exp[1]) < 0.01
                     and abs(got[2] - exp[2]) < 1e-6 and got[3] == exp[3]
                     and got[4].startswith(exp[4]) and not errors)
            ok &= match
            print(f"{'PASS' if match else 'FAIL'}  {name}")
            print(f"      cena ${got[0]:,.0f} | koszt ${got[1]:,.2f} | marża {got[2]:.0%} | "
                  f"cel ${got[3]:,.0f} | {got[4]}")
            if not match:
                print(f"      oczekiwano: {exp}  błędy formuł: {errors}")
            if i == 0:
                print("      COPY:", " / ".join(c.value for c in wb["COPY"]["A"][2:] if c.value))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
